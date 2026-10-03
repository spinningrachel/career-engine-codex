import hashlib
import importlib.util
import io
import json
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
SCRIPTS=ROOT/'plugin/skills/career-engine-export/scripts'

def load(name,path):
    spec=importlib.util.spec_from_file_location(name,path)
    module=importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
    return module

class ChatGPTTests(unittest.TestCase):
    def setUp(self):
        self.builder=load('chatgpt_builder',ROOT/'tools/build-chatgpt.py')

    def test_kit_is_reproducible_and_both_routes_preserve_all_source_sections(self):
        with tempfile.TemporaryDirectory() as tmp:
            archive=Path(tmp)/'kit.zip'
            self.builder.build(archive)
            self.assertEqual(archive.read_bytes(),(ROOT/'career-engine-chatgpt.zip').read_bytes())
        with zipfile.ZipFile(ROOT/'career-engine-chatgpt.zip') as kit:
            manifest=json.loads(kit.read('manifest.json'))
            self.assertEqual(manifest['installation_routes'],['custom-gpt','project'])
            self.assertEqual(manifest['upstream_commit'],json.loads((ROOT/'upstream-lock.json').read_text())['commit'])
            self.assertEqual(kit.read('UPDATE-CONTRACT.md'),(ROOT/'plugin/UPDATE-CONTRACT.md').read_bytes())
            for name,sha in manifest['file_sha256'].items():
                self.assertEqual(hashlib.sha256(kit.read(name)).hexdigest(),sha,name)
            source_paths={p.relative_to(ROOT/'plugin').as_posix() for folder in ('skills','agents','references') for p in (ROOT/'plugin'/folder).rglob('*') if p.is_file() and p.suffix in ('.md','.json','.csv') and not (p.name=='SKILL.md' and p.parent.name.startswith('role-') and p.parent.name!='role-prioritizer')}
            self.assertEqual(set(manifest['source_index']),source_paths)
            for path,info in manifest['source_index'].items():
                corpus=kit.read(info['knowledge_file']).decode()
                section=re.search(r'^=== SOURCE: '+re.escape(path)+r' ===\n(.*?)\n=== END SOURCE ===$',corpus,re.M|re.S)
                self.assertIsNotNone(section,path)
                self.assertEqual(section.group(1),self.builder.doctrine()[path].rstrip(),path)
            for mode in ('GPT','PROJECT'):
                instructions=kit.read(mode+'-INSTRUCTIONS.txt').decode()
                self.assertLessEqual(len(instructions),8000)
                self.assertIn('INSTALLATION MODE:',instructions)
            self.assertEqual(len([n for n in kit.namelist() if n.startswith('knowledge/')]),3)
            for name in ('detailed-cv','brief-cv','cover-letter'):
                doc=Document(io.BytesIO(kit.read('templates/'+name+'.docx')))
                self.assertNotIn('Homer',str(doc._element.xml))
                self.assertNotIn('Vandelay',str(doc._element.xml))

    def runtime(self,root):
        shutil.copy2(ROOT/'chatgpt/chatgpt_export.py',root/'chatgpt_export.py')
        shutil.copy2(SCRIPTS/'assemble_brief_cv.py',root/'assemble_brief_cv.py')
        return load('sandbox_export',root/'chatgpt_export.py')

    def test_portable_brief_export_preserves_roles_styles_and_photo_without_subprocess(self):
        sys.path.insert(0,str(SCRIPTS))
        try:
            fixture=load('upstream_fixture',SCRIPTS/'test_assemble_brief_cv.py')
        finally:
            sys.path.remove(str(SCRIPTS))
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); runtime=self.runtime(root)
            template=fixture.build_fixture_template_with_photo(tmp)
            output=root/'brief.docx'
            markdown=(SCRIPTS/'test-fixtures/brief-cv-sample.md').read_text()
            footer=(SCRIPTS/'test-fixtures/brief-footer-sample.md').read_text()
            # Compare parsed content with the real pandoc route, allowing smart
            # typography rather than changing facts or dropping sections.
            baseline=fixture.m.parse_cv_markdown(markdown)
            portable=runtime.assembler().parse_cv_markdown(markdown)
            normalize=lambda data: json.dumps(data,ensure_ascii=False).replace('’',"'").replace('–','--')
            self.assertEqual(normalize(baseline),normalize(portable))
            with patch.object(subprocess,'run',side_effect=AssertionError('ChatGPT must not call pandoc/shell')):
                runtime.brief_export(markdown,footer,template,output,name='Synthetic User',tagline='Fixture Role',contact=['Fixture City','fixture@example.com'])
            fixture.assert_structure(str(output),len(portable['roles']))
            self.assertEqual(fixture._blip_count(str(output)),fixture._blip_count(template))
            doc=Document(output); table=doc.tables[0]
            self.assertEqual(table.cell(0,1).paragraphs[0].text,'Synthetic User')
            self.assertEqual([p.style.name for p in table.cell(0,1).paragraphs],['Heading 1','Subtitle'])
            text='\n'.join(p.text for row in table.rows for cell in row.cells for p in cell.paragraphs)
            for expected in ('40%','zero downtime','Earlier:','University of Example','English','Spanish'):
                self.assertIn(expected,text)
            runs=[run for row in table.rows for cell in row.cells for p in cell.paragraphs for run in p.runs]
            self.assertTrue(any(r.text=="Newman's" and r.style.name=='ColorEmphasis' for r in runs))

    def test_detailed_and_letter_exports_use_blank_packaged_templates_and_separate_files(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); runtime=self.runtime(root)
            with zipfile.ZipFile(ROOT/'career-engine-chatgpt.zip') as kit:
                for stem in ('detailed-cv','cover-letter'):
                    (root/(stem+'.docx')).write_bytes(kit.read('templates/'+stem+'.docx'))
            cv=root/'cv.docx'; letter=root/'letter.docx'
            with patch.object(subprocess,'run',side_effect=AssertionError('Unexpected subprocess')):
                runtime.detailed_export('## SUMMARY\n\n**Synthetic** CV facts.\n\n::: {custom-style="RoleTitle"}\nFixture role\n:::\n\n- Preserved achievement.',root/'detailed-cv.docx',cv,identity={'name':'Synthetic User','contact':'fixture@example.com','tagline':'Fixture Role'},tagline='Fixture Role')
                runtime.detailed_export('Dear fixture team,\n\nA separate letter with *emphasis*.\n\nSynthetic User',root/'cover-letter.docx',letter)
            doc=Document(cv)
            self.assertEqual(doc.paragraphs[0].style.name,'Heading 2')
            self.assertTrue(any(r.bold and r.text=='Synthetic' for p in doc.paragraphs for r in p.runs))
            self.assertTrue(any(p.style.name=='RoleTitle' and p.text=='Fixture role' for p in doc.paragraphs))
            self.assertIn('Synthetic User','\n'.join(p.text for s in doc.sections for p in s.first_page_header.paragraphs))
            self.assertNotIn('{{',doc._element.xml+doc.sections[0].first_page_header._element.xml)
            self.assertIn('separate letter','\n'.join(p.text for p in Document(letter).paragraphs))
            self.assertNotIn('separate letter','\n'.join(p.text for p in doc.paragraphs))

    def test_unsupported_markdown_and_missing_identity_fail_before_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); runtime=self.runtime(root)
            template=root/'template.docx'; doc=Document(); doc.sections[0].header.paragraphs[0].text='{{name}}'; doc.save(template)
            output=root/'out.docx'
            for markdown in ('| data | lost |','[link](https://example.invalid)','```python\npass\n```','::: {custom-style="Normal"}\nunclosed','1. numbered item'):
                with self.subTest(markdown=markdown),self.assertRaises(ValueError):
                    runtime.detailed_export(markdown,template,output,identity={'name':'Fixture'})
                self.assertFalse(output.exists())
            with self.assertRaisesRegex(ValueError,'Unresolved template identity'):
                runtime.detailed_export('Plain text',template,output)
            self.assertFalse(output.exists())
            with self.assertRaises(ValueError):
                runtime.brief_export('## EXPERIENCE\n\n::: {custom-style="UnknownStyle"}\nWould disappear\n:::',None,template,output,name='Fixture',tagline='Role',contact=[])
            self.assertFalse(output.exists())

    def test_compatibility_checker_rejects_stale_chatgpt_kit(self):
        checker=load('chatgpt_checker',ROOT/'tools/check-chatgpt-compatibility.py')
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            for path in ('tools','plugin','chatgpt','codex'):
                (root/path).symlink_to(ROOT/path,target_is_directory=True)
            for path in ('LICENSE','upstream-lock.json'):
                (root/path).symlink_to(ROOT/path)
            (root/'career-engine-chatgpt.zip').write_bytes((ROOT/'career-engine-chatgpt.zip').read_bytes())
            with zipfile.ZipFile(root/'career-engine-chatgpt.zip','a') as kit:
                kit.writestr('unexpected.txt','Unreviewed material')
            checker.ROOT=root
            with self.assertRaisesRegex(ValueError,'stale, incomplete'):
                checker.check()

if __name__=='__main__':
    unittest.main()
