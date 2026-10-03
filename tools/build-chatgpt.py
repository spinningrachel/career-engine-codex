#!/usr/bin/env python3
"""Build a reproducible installation kit for Custom GPTs and ChatGPT Projects."""
import argparse
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import re
import zipfile
from docx import Document

ROOT=Path(__file__).resolve().parents[1]

def doctrine():
    result={}
    for folder in ('skills','agents','references'):
        for path in sorted((ROOT/'plugin'/folder).rglob('*')):
            if not path.is_file() or path.suffix not in ('.md','.json','.csv'):
                continue
            rel=path.relative_to(ROOT/'plugin').as_posix()
            if path.name=='SKILL.md' and path.parent.name.startswith('role-') and path.parent.name!='role-prioritizer':
                continue  # Codex wrapper pointers; complete agent doctrine is indexed below.
            text=path.read_text()
            text=re.sub(r'(?m)^> \*\*Codex (host|QA entrypoint):\*\*.*\n','',text)
            text=text.replace('${CAREER_ENGINE_ROOT}/','').replace('$CAREER_ENGINE_ROOT/','')
            result[rel]=text
    return result

def docx_template(blob):
    # Remove fictional sample identities/content while retaining template styles,
    # section settings, table geometry, and drawing relationships.
    converted=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(blob)) as source,zipfile.ZipFile(converted,'w') as target:
        for name in source.namelist():
            data=source.read(name)
            if name=='[Content_Types].xml':
                data=data.replace(b'application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml',b'application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml')
            target.writestr(name,data)
    doc=Document(converted)
    placeholders={'Heading 1':'{{name}}','Subtitle':'{{tagline}}','PersonalDetails':'{{contact}}'}
    for part in doc.part.package.parts:
        if part.partname.startswith('/word/') and hasattr(part,'element'):
            for paragraph in part.element.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'):
                from docx.text.paragraph import Paragraph
                p=Paragraph(paragraph,doc)
                texts=paragraph.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')
                if texts:
                    texts[0].text=placeholders.get(p.style.name,'')
                    for node in texts[1:]:
                        node.text=''
    clean=io.BytesIO(); doc.save(clean)
    output=io.BytesIO()
    with zipfile.ZipFile(clean) as source,zipfile.ZipFile(output,'w',compression=zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            data=source.read(name)
            if name=='[Content_Types].xml':
                data=data.replace(b'application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml',b'application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml')
            info=zipfile.ZipInfo(name,date_time=(2026,1,1,0,0,0)); info.compress_type=zipfile.ZIP_DEFLATED
            target.writestr(info,data)
    return output.getvalue()

def files():
    sources=doctrine()
    groups={'knowledge/01-workflows-and-roles.txt':[],'knowledge/02-reference-templates.txt':[]}
    index={}
    for path,text in sources.items():
        group='knowledge/02-reference-templates.txt' if path.startswith('references/') else 'knowledge/01-workflows-and-roles.txt'
        groups[group].append('=== SOURCE: '+path+' ===\n'+text.rstrip()+'\n=== END SOURCE ===\n')
        index[path]={'knowledge_file':group,'sha256':hashlib.sha256(text.encode()).hexdigest()}
    guide=(ROOT/'chatgpt/install.md').read_bytes()
    contract=(ROOT/'codex/update-contract.md').read_bytes()
    runtime=(ROOT/'chatgpt/runtime.md').read_text()
    entries={'START-HERE.md':guide,'UPDATE-CONTRACT.md':contract}
    common=(ROOT/'chatgpt/instructions.txt').read_text()
    for mode in ('GPT','PROJECT'):
        tail='\nINSTALLATION MODE: '+('Custom GPT: use GPT Instructions and public Knowledge. Configure optional Actions only in its editor.' if mode=='GPT' else 'ChatGPT Project: use Project instructions and public Project files. Use only apps/tools available inside this Project; do not try to import GPT Actions.')+'\n'
        instructions=common+tail
        if len(instructions)>8000:
            raise ValueError('ChatGPT instructions exceed the conservative 8000-character budget')
        entries[mode+'-INSTRUCTIONS.txt']=instructions.encode()
    for name,parts in groups.items():
        entries[name]='\n'.join(parts).encode()
    index_text='\n'.join(path+' -> '+item['knowledge_file'] for path,item in index.items())
    entries['knowledge/00-host-and-index.txt']=(runtime+'\n\n'+contract.decode()+'\n\n## Source index\n'+index_text+'\n').encode()
    entries['runtime/chatgpt_export.py.txt']=(ROOT/'chatgpt/chatgpt_export.py').read_bytes()
    entries['runtime/assemble_brief_cv.py.txt']=(ROOT/'plugin/skills/career-engine-export/scripts/assemble_brief_cv.py').read_bytes()
    entries['actions/notion-openapi.json']=(ROOT/'chatgpt/notion-openapi.json').read_bytes()
    for stem,name in [('cv-template-default','detailed-cv'),('cv-template-brief-default','brief-cv'),('cover-letter-template','cover-letter')]:
        entries['templates/'+name+'.docx']=docx_template((ROOT/'plugin/references'/f'{stem}.dotx').read_bytes())
    lock=json.loads((ROOT/'upstream-lock.json').read_text())
    entries['LICENSE']=(ROOT/'LICENSE').read_bytes()
    manifest={'format':'chatgpt-installation-kit','installation_routes':['custom-gpt','project'],'upstream_repository':lock['repository'],'upstream_commit':lock['commit'],'codex_plugin_version':json.loads((ROOT/'plugin/.codex-plugin/plugin.json').read_text())['version'],'source_index':index,'file_sha256':{name:hashlib.sha256(data).hexdigest() for name,data in sorted(entries.items())},'validation_scope':'Offline packaging/import constraints and Python exports; live ChatGPT UI, enabled tools, and authenticated Notion operations require account-level verification.'}
    entries['manifest.json']=(json.dumps(manifest,indent=2)+'\n').encode()
    spec=importlib.util.spec_from_file_location('pdd',ROOT/'plugin/scripts/personal_data_detect.py')
    pdd=importlib.util.module_from_spec(spec); spec.loader.exec_module(pdd)
    for path,text in sources.items():
        if list(pdd.scan_text(text,path,max_hits=3)):
            raise ValueError('Personal data in ChatGPT doctrine: '+path)
    for name,data in entries.items():
        if name.startswith('knowledge/') and len(data)>2*1024*1024:
            raise ValueError('Knowledge file exceeds the package upload budget: '+name)
        if name.endswith('.docx'):
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                text='\n'.join(re.sub('<[^>]+>',' ',z.read(n).decode()) for n in z.namelist() if n.endswith('.xml'))
        else:
            text=data.decode('utf-8')
        if list(pdd.scan_text(text,name,max_hits=3)):
            raise ValueError('Personal data in ChatGPT package: '+name)
    return entries

def build(output):
    entries=files()
    output.parent.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(output,'w',compression=zipfile.ZIP_DEFLATED) as z:
        for name,data in sorted(entries.items()):
            info=zipfile.ZipInfo(name,date_time=(2026,1,1,0,0,0))
            info.compress_type=zipfile.ZIP_DEFLATED; info.external_attr=0o644<<16
            z.writestr(info,data)
    print('Built',output,'for Custom GPT and Project:',len(entries),'files; data scan passed')

if __name__=='__main__':
    ap=argparse.ArgumentParser()
    ap.add_argument('--output',type=Path,default=ROOT/'career-engine-chatgpt.zip')
    build(ap.parse_args().output)
